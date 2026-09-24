from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
from app.services.settings_service import SettingsService
from app.helpers.date_helper import DateHelper

class ReportExportHelper:

    @staticmethod
    def export_table_to_excel(
        table,
        file_path,
        report_title,
        from_date=None,
        to_date=None,
        filters=None,
        summary=None,
    ):

        currency_symbol = (
            SettingsService.currency_symbol()
        )
        
        workbook = Workbook()

        worksheet = workbook.active
        worksheet.title = "Report"

        column_count = table.columnCount()
        row_count = table.rowCount()

        # -------------------------------------------------
        # Report title
        # -------------------------------------------------

        if column_count > 0:
            worksheet.merge_cells(
                start_row=1,
                start_column=1,
                end_row=1,
                end_column=column_count
            )

        title_cell = worksheet.cell(
            row=1,
            column=1,
            value=report_title
        )

        title_cell.font = Font(
            bold=True,
            size=16
        )

        title_cell.alignment = Alignment(
            horizontal="center"
        )

        # -------------------------------------------------
        # Report period and filters
        # -------------------------------------------------

        current_row = 2

        date_based_reports = {
            "Work Orders",
            "Maintenance Costs",
            "Technician Performance",
            "Inventory Transactions",
            "Asset Maintenance History",
            "Technician Work History",
             "Purchase Orders",
        }

        if (
            report_title in date_based_reports
            and from_date
            and to_date
        ):
            period_cell = worksheet.cell(
                row=current_row,
                column=1,
                value=(
                    f"Period: "
                    f"{DateHelper.display(from_date)} "
                    f"to "
                    f"{DateHelper.display(to_date)}"
                )
            )

            period_cell.font = Font(
                italic=True,
                size=10
            )

            current_row += 1

        # -------------------------------------------------
        # Active filters
        # -------------------------------------------------

        if filters:

            active_filters = []

            for name, value in filters.items():

                if value:
                    active_filters.append(
                        f"{name}: {value}"
                    )

            if active_filters:

                filter_cell = worksheet.cell(
                    row=current_row,
                    column=1,
                    value=" | ".join(active_filters)
                )

                filter_cell.font = Font(
                    italic=True,
                    size=10
                )

                current_row += 1

        # -------------------------------------------------
        # Report summary
        # -------------------------------------------------

        summary_start_row = current_row

        if summary:

            summary_items = [
                (label, value)
                for label, value in summary.items()
                if label
            ]

            for column, (label, value) in enumerate(
                summary_items,
                start=1
            ):

                label_cell = worksheet.cell(
                    row=summary_start_row,
                    column=column,
                    value=label
                )

                label_cell.font = Font(
                    bold=True
                )

                label_cell.alignment = Alignment(
                    horizontal="center"
                )

                value_cell = worksheet.cell(
                    row=summary_start_row + 1,
                    column=column,
                    value=value
                )

                value_cell.alignment = Alignment(
                    horizontal="center"
                )

            header_row = (
                summary_start_row + 3
            )

        else:

            header_row = 3

        # -------------------------------------------------
        # Column headings
        # -------------------------------------------------

        for column in range(column_count):

            header_item = (
                table.horizontalHeaderItem(
                    column
                )
            )

            heading = (
                header_item.text()
                if header_item is not None
                else ""
            )

            cell = worksheet.cell(
                row=header_row,
                column=column + 1,
                value=heading
            )

            cell.font = Font(
                bold=True
            )

            cell.alignment = Alignment(
                horizontal="center"
            )

        # -------------------------------------------------
        # Report data
        # -------------------------------------------------

        for row in range(row_count):

            for column in range(column_count):

                item = table.item(
                    row,
                    column
                )

                text = (
                    item.text().strip()
                    if item is not None
                    else ""
                )

                cell = worksheet.cell(
                    row=header_row + row + 1,
                    column=column + 1
                )

                # -----------------------------------------
                # Empty cell
                # -----------------------------------------

                if text == "":
                    cell.value = ""
                    continue

                # -----------------------------------------
                # Currency value
                # -----------------------------------------

                if text.startswith(currency_symbol):
                    try:
                        numeric_value = float(
                            text
                            .replace(currency_symbol, "")
                            .replace(",", "")
                            .strip()
                        )

                        cell.value = numeric_value
                        cell.number_format = (
                            f'"{currency_symbol}" #,##0.00'
                        )
                    except ValueError:
                        cell.value = text
                    continue

                # -----------------------------------------
                # Normal numeric value
                # -----------------------------------------

                try:
                    numeric_value = float(
                        text.replace(",", "")
                    )

                    cell.value = numeric_value

                    header_item = (
                        table.horizontalHeaderItem(
                            column
                        )
                    )

                    heading = (
                        header_item.text().strip()
                        if header_item is not None
                        else ""
                    )

                    # -----------------------------------------
                    # Quantity change
                    # -----------------------------------------

                    if heading == "Change":

                        if "." in text:
                            cell.number_format = (
                                "+#,##0.00;-#,##0.00;0.00"
                            )
                        else:
                            cell.number_format = (
                                "+#,##0;-#,##0;0"
                            )

                    # -----------------------------------------
                    # Other numeric values
                    # -----------------------------------------

                    elif "." in text:

                        cell.number_format = (
                            "#,##0.00"
                        )

                except ValueError:
                    cell.value = text

        # -------------------------------------------------
        # Automatically size columns
        # -------------------------------------------------

        for column in range(
            1,
            column_count + 1
        ):
            max_length = 0

            column_letter = get_column_letter(
                column
            )

            header_cell = worksheet.cell(
                row=header_row,
                column=column
            )

            heading = str(
                header_cell.value or ""
            ).strip()

            for cell in worksheet[
                column_letter
            ]:
                if cell.value is None:
                    continue

                max_length = max(
                    max_length,
                    len(str(cell.value))
                )

            # Notes needs extra width and wrapping
            if heading == "Notes":
                worksheet.column_dimensions[
                    column_letter
                ].width = 55

                for row_number in range(
                    header_row + 1,
                    header_row + row_count + 1
                ):
                    worksheet.cell(
                        row=row_number,
                        column=column
                    ).alignment = Alignment(
                        vertical="top",
                        wrap_text=True
                    )
            else:

                column_width = min(
                    max_length + 2,
                    40
                )

                # Currency values display wider than their
                # underlying numeric values in Excel.
                currency_headings = {
                    "Estimated Cost",
                    "Actual Cost",
                    "Labour Cost",
                    "Material Cost",
                    "Total Cost",
                    "Unit Cost",
                    "Stock Value",
                    "Reorder Cost",
                    "Total",
                }

                if heading in currency_headings:
                    column_width = max(
                        column_width,
                        14
                    )

                worksheet.column_dimensions[
                    column_letter
                ].width = column_width
                
        # -------------------------------------------------
        # Freeze headings
        # -------------------------------------------------

        worksheet.freeze_panes = (
            f"A{header_row + 1}"
        )

        # -------------------------------------------------
        # Excel filters
        # -------------------------------------------------

        if (
            column_count > 0
            and row_count > 0
        ):
            worksheet.auto_filter.ref = (
                f"A{header_row}:"
                f"{get_column_letter(column_count)}"
                f"{header_row + row_count}"
            )

        # -------------------------------------------------
        # Save workbook
        # -------------------------------------------------

        workbook.save(
            file_path
        )